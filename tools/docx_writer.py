"""
OLG Document Writer — converts markdown drafts to styled .docx files.

Uses the OLG template .docx for styles when available, falls back to
basic formatting without one.

OLG Template Styles (2026):
  - Heading 1: Segoe UI Semibold 21pt, bold, #141414
  - Heading 2: Segoe UI Semibold 15pt, bold, #14558F (blue)
  - Heading 3: Segoe UI Semibold 12.5pt, bold italic, #388557 (green)
  - Heading 4: bold italic, #365F91 (steel blue)
  - Normal: #141414
  - Note: callout style (label bolded, e.g. "Note:" or "Important:")

Usage:
    from tools.docx_writer import generate_docx
    generate_docx(markdown_text, output_path, template_path)
"""

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# OLG text color
OLG_TEXT_COLOR = RGBColor(0x14, 0x14, 0x14)


# ============================================================
# MARKDOWN PARSER
# ============================================================

def _parse_markdown(text: str) -> list[dict]:
    """Parse markdown into a list of block elements.

    Returns a list of dicts, each with:
        type: heading | paragraph | bullet | numbered | table | separator | note
        level: heading level (1-6)
        text: the text content (with inline markdown preserved)
        label: for notes, the label text (e.g. "Note:" or "Important:")
        rows: for tables, list of row lists
    """
    blocks = []
    lines = text.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            i += 1
            continue

        # Separator
        if stripped == "---" or stripped == "***" or stripped == "___":
            blocks.append({"type": "separator"})
            i += 1
            continue

        # Heading
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            blocks.append({
                "type": "heading",
                "level": level,
                "text": heading_match.group(2).strip(),
            })
            i += 1
            continue

        # Table (pipe-delimited)
        if "|" in stripped and stripped.startswith("|"):
            rows = []
            while i < len(lines) and "|" in lines[i].strip():
                row_line = lines[i].strip()
                # Skip separator rows (|---|---|)
                if re.match(r"^\|[\s\-:|]+\|$", row_line):
                    i += 1
                    continue
                cells = [c.strip() for c in row_line.split("|")[1:-1]]
                rows.append(cells)
                i += 1
            if rows:
                blocks.append({"type": "table", "rows": rows})
            continue

        # Bullet list
        bullet_match = re.match(r"^[*\-+]\s+(.+)$", stripped)
        if bullet_match:
            items = []
            while i < len(lines):
                bm = re.match(r"^[*\-+]\s+(.+)$", lines[i].strip())
                if not bm:
                    break
                items.append(bm.group(1))
                i += 1
            for item in items:
                blocks.append({"type": "bullet", "text": item})
            continue

        # Numbered list
        num_match = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if num_match:
            items = []
            while i < len(lines):
                nm = re.match(r"^\d+[.)]\s+(.+)$", lines[i].strip())
                if not nm:
                    break
                items.append(nm.group(1))
                i += 1
            for item in items:
                blocks.append({"type": "numbered", "text": item})
            continue

        # Blockquote
        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip("> "))
                i += 1
            blocks.append({"type": "blockquote", "text": " ".join(quote_lines)})
            continue

        # Note/Important callout — detect "Note:" or "Important:" at start
        note_match = re.match(
            r"^(Note|Important|Warning|Reminder|Tip):\s*(.*)$",
            stripped, re.IGNORECASE,
        )
        if note_match:
            label = note_match.group(1) + ":"
            body_parts = [note_match.group(2)]
            i += 1
            # Collect continuation lines
            while i < len(lines):
                next_line = lines[i].strip()
                if not next_line:
                    break
                if (re.match(r"^#{1,6}\s+", next_line) or
                    re.match(r"^[*\-+]\s+", next_line) or
                    re.match(r"^\d+[.)]\s+", next_line) or
                    next_line.startswith("|") or
                    next_line in ("---", "***", "___") or
                    next_line.startswith(">")):
                    break
                body_parts.append(next_line)
                i += 1
            blocks.append({
                "type": "note",
                "label": label,
                "text": " ".join(body_parts),
            })
            continue

        # Regular paragraph — collect continuation lines
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            next_line = lines[i].strip()
            if not next_line:
                break
            if (re.match(r"^#{1,6}\s+", next_line) or
                re.match(r"^[*\-+]\s+", next_line) or
                re.match(r"^\d+[.)]\s+", next_line) or
                next_line.startswith("|") or
                next_line in ("---", "***", "___") or
                next_line.startswith(">") or
                re.match(r"^(Note|Important|Warning|Reminder|Tip):",
                         next_line, re.IGNORECASE)):
                break
            para_lines.append(next_line)
            i += 1

        blocks.append({"type": "paragraph", "text": " ".join(para_lines)})

    return blocks


# ============================================================
# INLINE FORMATTING
# ============================================================

def _apply_inline_formatting(paragraph, text: str, default_color=None):
    """Apply bold, italic, and other inline markdown to a paragraph.

    Handles **bold**, *italic*, ***bold italic***, and `code`.
    """
    pattern = re.compile(
        r"(\*\*\*.*?\*\*\*)"   # ***bold italic***
        r"|(\*\*.*?\*\*)"      # **bold**
        r"|(\*.*?\*)"          # *italic*
        r"|(`[^`]+`)"          # `code`
    )

    parts = pattern.split(text)

    for part in parts:
        if part is None or part == "":
            continue

        if part.startswith("***") and part.endswith("***"):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(part)

        if default_color and run.font.color.rgb is None:
            run.font.color.rgb = default_color


# ============================================================
# STYLE DETECTION
# ============================================================

def _get_available_styles(doc: Document) -> dict[str, str]:
    """Detect which styles are available in the document.

    Returns a mapping from our element types to the actual style names.
    Prioritizes OLG-specific styles from the template.
    """
    available = {s.name for s in doc.styles}

    style_map = {}

    # Headings
    for level in range(1, 7):
        if f"Heading {level}" in available:
            style_map[f"heading_{level}"] = f"Heading {level}"

    # Body text
    if "Normal" in available:
        style_map["paragraph"] = "Normal"

    # Note/callout — OLG template has a custom "Note" style
    if "Note" in available:
        style_map["note"] = "Note"

    # Bullet lists
    for candidate in ["List Bullet", "List Bullet 2"]:
        if candidate in available:
            style_map["bullet"] = candidate
            break

    # Numbered lists
    for candidate in ["List Number", "List Number 2"]:
        if candidate in available:
            style_map["numbered"] = candidate
            break

    # Blockquote
    for candidate in ["Note", "Block Text", "Quote", "Body Text Indent"]:
        if candidate in available:
            style_map["blockquote"] = candidate
            break

    return style_map


# ============================================================
# LIST NUMBERING (OLG TEMPLATE)
# ============================================================

class _ListNumbering:
    """Manages list numbering for OLG template documents.

    The OLG template uses Normal-style paragraphs with XML numbering
    properties (w:numPr) rather than dedicated list styles. This class
    creates new numbering instances for each list in the output.
    """

    def __init__(self, doc: Document):
        self._doc = doc
        self._next_num_id = 100  # Start high to avoid collisions
        self._bullet_abstract_id = None
        self._decimal_abstract_id = None
        self._current_bullet_num_id = None
        self._current_decimal_num_id = None
        self._detect_abstract_nums()

    def _detect_abstract_nums(self):
        """Find bullet and decimal abstractNum IDs from the template."""
        try:
            numbering_part = self._doc.part.numbering_part
            numbering_xml = numbering_part._element
        except Exception:
            return

        for abstract in numbering_xml.findall(qn('w:abstractNum')):
            abs_id = abstract.get(qn('w:abstractNumId'))
            lvl0 = abstract.find(qn('w:lvl'))
            if lvl0 is not None:
                numFmt = lvl0.find(qn('w:numFmt'))
                fmt = numFmt.get(qn('w:val')) if numFmt is not None else ''
                if fmt == 'bullet' and self._bullet_abstract_id is None:
                    self._bullet_abstract_id = abs_id
                elif fmt == 'decimal' and self._decimal_abstract_id is None:
                    self._decimal_abstract_id = abs_id

    def _create_num_id(self, abstract_id: str) -> int:
        """Create a new w:num entry referencing an abstractNum."""
        try:
            numbering_part = self._doc.part.numbering_part
            numbering_xml = numbering_part._element
        except Exception:
            return 0

        num_id = self._next_num_id
        self._next_num_id += 1

        num_elem = OxmlElement('w:num')
        num_elem.set(qn('w:numId'), str(num_id))
        abs_ref = OxmlElement('w:abstractNumId')
        abs_ref.set(qn('w:val'), abstract_id)
        num_elem.append(abs_ref)
        numbering_xml.append(num_elem)

        return num_id

    def start_bullet_list(self):
        """Start a new bullet list (resets numbering)."""
        if self._bullet_abstract_id:
            self._current_bullet_num_id = self._create_num_id(
                self._bullet_abstract_id)

    def start_numbered_list(self):
        """Start a new numbered list (resets numbering)."""
        if self._decimal_abstract_id:
            self._current_decimal_num_id = self._create_num_id(
                self._decimal_abstract_id)

    def apply_bullet(self, paragraph):
        """Apply bullet numbering to a paragraph."""
        if self._current_bullet_num_id:
            self._set_numPr(paragraph, self._current_bullet_num_id, 0)

    def apply_numbered(self, paragraph):
        """Apply numbered list numbering to a paragraph."""
        if self._current_decimal_num_id:
            self._set_numPr(paragraph, self._current_decimal_num_id, 0)

    def _set_numPr(self, paragraph, num_id: int, ilvl: int):
        """Set w:numPr on a paragraph's properties."""
        pPr = paragraph._p.get_or_add_pPr()
        numPr = OxmlElement('w:numPr')
        ilvl_elem = OxmlElement('w:ilvl')
        ilvl_elem.set(qn('w:val'), str(ilvl))
        numId_elem = OxmlElement('w:numId')
        numId_elem.set(qn('w:val'), str(num_id))
        numPr.append(ilvl_elem)
        numPr.append(numId_elem)
        pPr.append(numPr)

    @property
    def has_numbering(self) -> bool:
        return (self._bullet_abstract_id is not None or
                self._decimal_abstract_id is not None)


# ============================================================
# DOCUMENT GENERATION
# ============================================================

def generate_docx(
    markdown_text: str,
    output_path: Path | str,
    template_path: Path | str | None = None,
) -> Path:
    """Convert markdown text to a styled .docx file.

    Args:
        markdown_text: The draft content in markdown format.
        output_path: Where to save the .docx file.
        template_path: Optional OLG template .docx to use for styles.
                       If provided, styles from the template are used.
                       If not, basic formatting is applied.

    Returns:
        The output path.
    """
    output_path = Path(output_path)

    if template_path and Path(template_path).exists():
        doc = Document(str(template_path))
        # Clear the template's body content but keep styles and section properties.
        # w:sectPr defines page size/margins and must be preserved — python-docx
        # needs it for operations like add_table() that compute block width.
        for element in doc.element.body[:]:
            if element.tag != qn('w:sectPr'):
                doc.element.body.remove(element)
        has_template = True
    else:
        doc = Document()
        has_template = False

    style_map = _get_available_styles(doc)
    blocks = _parse_markdown(markdown_text)
    list_numbering = _ListNumbering(doc) if has_template else None

    # Strip author's notes section (after --- separator near the end)
    author_notes_start = None
    for idx, block in enumerate(blocks):
        if block["type"] == "separator":
            if idx + 1 < len(blocks) and blocks[idx + 1]["type"] == "heading":
                heading_text = blocks[idx + 1]["text"].lower()
                if "author" in heading_text and "note" in heading_text:
                    author_notes_start = idx
                    break

    if author_notes_start is not None:
        blocks = blocks[:author_notes_start]

    # Default text color when using OLG template
    text_color = OLG_TEXT_COLOR if has_template else None

    # Track list continuity — start new numbering when list type changes
    prev_type = None

    for block in blocks:
        btype = block["type"]

        if btype == "heading":
            level = block["level"]
            style_key = f"heading_{level}"
            style_name = style_map.get(style_key)

            if style_name:
                para = doc.add_paragraph(style=style_name)
                _apply_inline_formatting(para, block["text"])
            else:
                doc.add_heading(block["text"], level=level)

        elif btype == "note":
            # OLG Note style: bold label + body text
            style_name = style_map.get("note")
            if style_name:
                para = doc.add_paragraph(style=style_name)
            else:
                para = doc.add_paragraph()

            label_run = para.add_run(block["label"])
            label_run.bold = True
            if text_color:
                label_run.font.color.rgb = text_color

            space_run = para.add_run(" ")
            if text_color:
                space_run.font.color.rgb = text_color

            _apply_inline_formatting(para, block["text"], text_color)

        elif btype == "paragraph":
            style_name = style_map.get("paragraph", "Normal")
            para = doc.add_paragraph(style=style_name)
            _apply_inline_formatting(para, block["text"], text_color)

        elif btype == "bullet":
            # Start new bullet list if previous block wasn't a bullet
            if prev_type != "bullet" and list_numbering and list_numbering.has_numbering:
                list_numbering.start_bullet_list()

            if list_numbering and list_numbering.has_numbering:
                para = doc.add_paragraph(style="Normal")
                list_numbering.apply_bullet(para)
            else:
                style_name = style_map.get("bullet", "List Bullet")
                try:
                    para = doc.add_paragraph(style=style_name)
                except KeyError:
                    para = doc.add_paragraph()
                    para.paragraph_format.left_indent = Inches(0.5)
            _apply_inline_formatting(para, block["text"], text_color)

        elif btype == "numbered":
            # Start new numbered list if previous block wasn't numbered
            if prev_type != "numbered" and list_numbering and list_numbering.has_numbering:
                list_numbering.start_numbered_list()

            if list_numbering and list_numbering.has_numbering:
                para = doc.add_paragraph(style="Normal")
                list_numbering.apply_numbered(para)
            else:
                style_name = style_map.get("numbered", "List Number")
                try:
                    para = doc.add_paragraph(style=style_name)
                except KeyError:
                    para = doc.add_paragraph()
                    para.paragraph_format.left_indent = Inches(0.5)
            _apply_inline_formatting(para, block["text"], text_color)

        elif btype == "blockquote":
            style_name = style_map.get("blockquote")
            if style_name:
                para = doc.add_paragraph(style=style_name)
            else:
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Inches(0.5)
                para.paragraph_format.right_indent = Inches(0.5)
            _apply_inline_formatting(para, block["text"], text_color)

        elif btype == "table":
            rows = block["rows"]
            if not rows:
                continue
            num_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = "Table Grid"

            for r_idx, row in enumerate(rows):
                for c_idx, cell_text in enumerate(row):
                    if c_idx < num_cols:
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        if r_idx == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True

        elif btype == "separator":
            doc.add_paragraph()

        prev_type = btype

    doc.save(str(output_path))
    return output_path
